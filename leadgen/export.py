"""Write the run's results to CSV and Excel for the business team."""
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, db

COLUMNS = [
    ("Doctor Name", "name"), ("Specialty", "specialty"), ("Hospital/Clinic", "hospitals"), ("Location", "location"),
    ("Email", "email"), ("Email Verification/Status", "email_status"), ("Email Type", "email_type"),
    ("Email Source URL", "email_source_url"), ("Source/Profile URL", "profile_urls"), ("MCHK Reg. No.", "mchk_reg_no"),
    ("Lead Status", "lead_status"), ("Reason Codes", "reason_codes"), ("Verification Notes", "notes"),
]

STATUS_ORDER = {"QUALIFIED": 0, "NEEDS_REVIEW": 1, "PENDING_ENRICHMENT": 2, "PENDING": 3, "EXISTING_IN_CRM": 4, "REJECTED": 5}


def _row(lead):
    out = []
    for _, key in COLUMNS:
        v = lead[key]
        if key in ("hospitals", "profile_urls", "reason_codes"):
            v = "; ".join(db.jload(v))
        elif key == "notes":
            v = " | ".join(db.jload(v))
        out.append(v or "")
    return out


def export_run(conn, run_id: int, target: int):
    leads = [dict(r) for r in conn.execute("SELECT * FROM leads WHERE run_id=?", (run_id,))]
    leads.sort(key=lambda l: (STATUS_ORDER.get(l["lead_status"], 9), l["name"]))
    qualified = [l for l in leads if l["lead_status"] == "QUALIFIED"][:target]
    config.OUTPUT_DIR.mkdir(exist_ok=True)
    paths = []

    for fname, rows in ((f"run{run_id}_qualified_leads.csv", qualified), (f"run{run_id}_all_leads.csv", leads)):
        p = config.OUTPUT_DIR / (config.EXPORT_PREFIX + fname)
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow([c for c, _ in COLUMNS])
            w.writerows(_row(l) for l in rows)
        paths.append(p)

    wb = Workbook()
    fills = {"QUALIFIED": "C6EFCE", "NEEDS_REVIEW": "FFEB9C", "REJECTED": "FFC7CE", "EXISTING_IN_CRM": "D9D9D9",
             "PENDING_ENRICHMENT": "DDEBF7", "PENDING": "F2F2F2"}
    for title, rows in (("Qualified leads", qualified), ("All candidates", leads)):
        ws = wb.active if title == "Qualified leads" else wb.create_sheet(title)
        ws.title = title
        ws.append([c for c, _ in COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for l in rows:
            ws.append(_row(l))
            status_cell = ws.cell(row=ws.max_row, column=[k for _, k in COLUMNS].index("lead_status") + 1)
            status_cell.fill = PatternFill("solid", fgColor=fills.get(l["lead_status"], "FFFFFF"))
        widths = [26, 26, 34, 11, 32, 14, 14, 40, 50, 12, 18, 30, 90]
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
        ws.freeze_panes = "B2"
    p = config.OUTPUT_DIR / f"{config.EXPORT_PREFIX}run{run_id}_leads.xlsx"
    wb.save(p)
    paths.append(p)
    return paths
